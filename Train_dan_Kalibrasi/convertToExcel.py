import sqlite3
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DB_PATH = "conveyor_monitoring.db"
OUTPUT_EXCEL = "dataset_conveyor_kalibrasi.xlsx"

def ekspor_db_ke_excel_rapi():
    print("Membaca data dari SQLite database...")
    conn = sqlite3.connect(DB_PATH)
    
    # Ambil data dari tabel sesi pengamatan
    # Anda juga bisa menyesuaikan query SQL ini sesuai kebutuhan
    query = """
    SELECT 
        session_id AS "ID Sesi",
        lokasi AS "Lokasi",
        n_detections AS "Jumlah Deteksi (N)",
        r_max AS "Luas Maks (R_max)",
        n_norm AS "N Ternormalisasi",
        c_avg AS "Confidence Rata-rata (C_avg)",
        score_s AS "Skor SAW (S)",
        category_system AS "Kategori Rule-Based",
        category_logreg AS "Kategori Logreg",
        '' AS "Label Teknisi (Ground Truth)"  -- Kolom kosong khusus diisi teknisi
    FROM sesi_pengamatan
    ORDER BY session_id DESC
    """
    
    df = pd.read_sql_query(query, conn)
    conn.close()

    # Ubah nilai NULL/None menjadi 0.0 untuk angka numerik agar rapi
    df["Jumlah Deteksi (N)"] = df["Jumlah Deteksi (N)"].fillna(0)
    df["Luas Maks (R_max)"] = df["Luas Maks (R_max)"].fillna(0.0)
    df["N Ternormalisasi"] = df["N Ternormalisasi"].fillna(0.0)
    df["Confidence Rata-rata (C_avg)"] = df["Confidence Rata-rata (C_avg)"].fillna(0.0)
    df["Skor SAW (S)"] = df["Skor SAW (S)"].fillna(0.0)

    # Ekspor ke Excel menggunakan openpyxl untuk styling
    with pd.ExcelWriter(OUTPUT_EXCEL, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Dataset Sesi", index=False)
        worksheet = writer.sheets["Dataset Sesi"]

        # Stylings
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Biru Tua
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell_font = Font(name="Calibri", size=11)
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        # Apply style ke Header
        for col_num, col_name in enumerate(df.columns, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        # Apply style ke Sel Data & Atur Lebar Kolom Otomatis
        for row in worksheet.iter_rows(min_row=2, max_row=len(df)+1, min_col=1, max_col=len(df.columns)):
            for cell in row:
                cell.font = cell_font
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        # Auto-fit lebar kolom
        for col in worksheet.columns:
            max_length = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            worksheet.column_dimensions[col_letter].width = max(max_length + 4, 12)

    print(f"Selesai! File Excel rapi berhasil disimpan ke: {OUTPUT_EXCEL}")

if __name__ == "__main__":
    ekspor_db_ke_excel_rapi()